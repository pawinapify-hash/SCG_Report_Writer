import base64
import io
import json
import os
import re

import requests
from dotenv import load_dotenv
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

load_dotenv(override=True)

API_KEY = os.getenv("OPENAI_API_KEY", "").strip()
MODEL = "gpt-4o"
PPTX_PATH = "SCG Monthly Report Automate Template July2.pptx"
OUTPUT_PATH = "extracted_posts.json"
SLIDES = [15, 16, 17]  # 1-indexed slide numbers

EMU_PER_IN = 914400
MAX_IMG_W = 512
TAG_WORDS = {"Environmental", "Social", "Governance"}
MARKER_WORDS = {"ลบโพสต์", "เสร็จเเล้วค่ะ", "เสร็จแล้วค่ะ"}
ENG_RE = re.compile(
    r"Engagement:\s*([\d,]+)\s*Like:\s*([\d,]+)\s*Comment:\s*([\d,]+)\s*Share:\s*([\d,]+)"
)


def shape_text(sh):
    if sh.has_text_frame:
        return sh.text_frame.text.strip()
    return ""


def is_placeholder(text):
    return "$" in text


def downscale(data, max_w=MAX_IMG_W):
    im = Image.open(io.BytesIO(data)).convert("RGB")
    if im.width > max_w:
        ratio = max_w / im.width
        im = im.resize((max_w, int(im.height * ratio)), Image.LANCZOS)
    buf = io.BytesIO()
    im.save(buf, format="JPEG", quality=85)
    return buf.getvalue()


def image_data_url(shape):
    blob = shape.image.blob
    ext = (shape.image.ext or "png").lower()
    mime = "image/png" if ext == "png" else "image/jpeg"
    small = downscale(blob)
    return "data:" + mime + ";base64," + base64.b64encode(small).decode()


def extract_posts(slide):
    eng_shapes = []
    text_shapes = []
    pictures = []

    for sh in slide.shapes:
        t = shape_text(sh)
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            pictures.append(sh)
        elif t.startswith("Engagement:"):
            eng_shapes.append(sh)
        else:
            text_shapes.append(sh)

    posts = []
    for eng in eng_shapes:
        left = eng.left / EMU_PER_IN
        top = eng.top / EMU_PER_IN

        m = ENG_RE.search(shape_text(eng))
        if not m:
            continue
        engagement = {
            "engagement": int(m.group(1).replace(",", "")),
            "like": int(m.group(2).replace(",", "")),
            "comment": int(m.group(3).replace(",", "")),
            "share": int(m.group(4).replace(",", "")),
        }

        tag = None
        caption_candidates = []
        for sh in text_shapes:
            sl = sh.left / EMU_PER_IN
            st = sh.top / EMU_PER_IN
            txt = shape_text(sh)
            if not txt:
                continue
            if abs(sl - left) > 0.15:
                continue
            if not (top - 0.95 <= st <= top - 0.55):
                continue
            if txt in TAG_WORDS:
                tag = txt
            elif txt in MARKER_WORDS:
                continue
            elif not is_placeholder(txt) and not txt.startswith("Engagement:"):
                caption_candidates.append(txt)

        caption = max(caption_candidates, key=len) if caption_candidates else ""

        image = None
        post_center = left + 0.5
        for sh in pictures:
            sl = sh.left / EMU_PER_IN
            st = sh.top / EMU_PER_IN
            sw = sh.width / EMU_PER_IN
            if not (top - 1.2 <= st <= top - 0.1):
                continue
            pic_center = sl + sw / 2
            if abs(pic_center - post_center) < 0.4:
                image = sh
                break

        posts.append({
            "caption": caption,
            "engagement": engagement,
            "tag": tag,
            "image": image,
            "_left": left,
            "_top": top,
        })

    posts.sort(key=lambda p: (p["_top"], p["_left"]))
    return posts


def describe_images(posts):
    image_parts = []
    idx_map = []
    for i, p in enumerate(posts):
        if p["image"] is not None:
            image_parts.append(
                {"type": "image_url", "image_url": {"url": image_data_url(p["image"])}}
            )
            idx_map.append(i)

    if not image_parts:
        return {}

    n = len(image_parts)
    prompt = (
        f"คุณจะได้รับรูปภาพโพสต์โซเชียลมีเดียจำนวน {n} รูปตามลำดับ: รูปที่ 1 ถึงรูปที่ {n}\n"
        "สำหรับแต่ละรูป ให้บรรยายสิ่งที่เห็นในภาพเป็นภาษาไทย จำนวน 3 ประโยค:\n"
        "ประโยคที่ 1: ฉากหรือเนื้อหาหลักของภาพ บุคคล สิ่งของ สี และการจัดวางองค์ประกอบ\n"
        "ประโยคที่ 2: ข้อความ พาดหัว หรือโลโก้ที่ปรากฏบนภาพ (อ่านข้อความภาษาไทยบนภาพแล้วเขียนออกมาเป็นภาษาไทย)\n"
        "ประโยคที่ 3: สไตล์กราฟิกและโทนหรืออารมณ์โดยรวมของโพสต์\n\n"
        f"ตอบกลับเป็น JSON array จำนวน {n} object โดยมี key: post, image_description เท่านั้น ไม่ต้องมีข้อความอื่น"
    )

    content = [{"type": "text", "text": prompt}] + image_parts
    resp = requests.post(
        "https://api.openai.com/v1/chat/completions",
        headers={"Authorization": f"Bearer {API_KEY}", "Content-Type": "application/json"},
        json={"model": MODEL, "messages": [{"role": "user", "content": content}],
              "temperature": 0.2, "max_tokens": 4000},
        timeout=300,
    )
    resp.raise_for_status()
    data = resp.json()
    raw = data["choices"][0]["message"]["content"]

    m = re.search(r"\[.*\]", raw, re.DOTALL)
    descriptions = json.loads(m.group(0)) if m else []

    result = {}
    for d in descriptions:
        batch_idx = d.get("post")
        if batch_idx is None:
            continue
        i = batch_idx - 1
        if 0 <= i < len(idx_map):
            result[idx_map[i]] = d.get("image_description")

    return result


def main():
    prs = Presentation(PPTX_PATH)
    results = []
    usage_total = {}

    for slide_no in SLIDES:
        slide = prs.slides[slide_no - 1]
        posts = extract_posts(slide)

        desc_map = describe_images(posts)

        out_posts = []
        for i, p in enumerate(posts):
            out_posts.append({
                "post": i + 1,
                "caption": p["caption"],
                "engagement": p["engagement"],
                "tag": p["tag"],
                "image_description": desc_map.get(i),
            })

        results.append({"slide": slide_no, "posts": out_posts})
        print(f"slide {slide_no}: {len(out_posts)} posts")

    out = {"results": results, "usage_total": usage_total}
    with open(OUTPUT_PATH, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=2)
    print("saved ->", OUTPUT_PATH)


if __name__ == "__main__":
    main()
