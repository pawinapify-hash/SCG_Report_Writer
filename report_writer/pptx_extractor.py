import base64
import io
import json
import re

import requests
from PIL import Image
from pptx.enum.shapes import MSO_SHAPE_TYPE

EMU_PER_IN = 914400
MAX_IMG_W = 512
TAG_WORDS = {
    "Global ESG",
    "Environmental",
    "Social",
    "Governance",
    "Operating Results",
    "Financial Stability",
    "Business Competitiveness & Corporate",
    "Stock / Debenture",
    "Thought Leadership",
    "Business Award",
    "SCGD",
    "Activity",
    "Issue/Crisis",
    "Political",
    "New thing creator",
    "AI Explorer",
    "Happy Workplace",
    "Agile Workplace",
    "Innovation",
}
MARKER_WORDS = {"ลบโพสต์", "เสร็จเเล้วค่ะ", "เสร็จแล้วค่ะ"}
ENG_RE = re.compile(
    r"Engagement:\s*([\d,]+)\s*Like:\s*([\d,]+)\s*Comment:\s*([\d,]+)\s*Share:\s*([\d,]+)"
)
PRI_RE = re.compile(
    r"PRI:\s*([\d,]+)\s*PRE:\s*([\d.]+)\s*Engagement:\s*([\d,]+)"
)


def shape_text(sh):
    if sh.has_text_frame:
        return sh.text_frame.text.strip()
    return ""


def is_placeholder(text):
    return "$" in text


def downscale(data, max_w=MAX_IMG_W):
    im = Image.open(io.BytesIO(data)).convert("RGB")
    if im.width > max_w:
        ratio = max_w / im.width
        im = im.resize((max_w, int(im.height * ratio)), Image.LANCZOS)
    buf = io.BytesIO()
    im.save(buf, format="JPEG", quality=85)
    return buf.getvalue()


def image_data_url(blob, ext="png"):
    mime = "image/png" if ext == "png" else "image/jpeg"
    small = downscale(blob)
    return "data:" + mime + ";base64," + base64.b64encode(small).decode()


def extract_posts(slide):
    metric_shapes = []
    text_shapes = []
    pictures = []

    for sh in slide.shapes:
        t = shape_text(sh)
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            pictures.append(sh)
        elif t.startswith("Engagement:") or t.startswith("PRI:"):
            metric_shapes.append(sh)
        else:
            text_shapes.append(sh)

    posts = []
    for met in metric_shapes:
        left = met.left / EMU_PER_IN
        top = met.top / EMU_PER_IN
        t = shape_text(met)

        is_pri = t.startswith("PRI:")
        if is_pri:
            m = PRI_RE.search(t)
            if not m:
                continue
            metrics = {
                "pri": int(m.group(1).replace(",", "")),
                "pre": float(m.group(2)),
                "engagement": int(m.group(3).replace(",", "")),
            }
        else:
            m = ENG_RE.search(t)
            if not m:
                continue
            metrics = {
                "engagement": int(m.group(1).replace(",", "")),
                "like": int(m.group(2).replace(",", "")),
                "comment": int(m.group(3).replace(",", "")),
                "share": int(m.group(4).replace(",", "")),
            }

        tag = None
        publisher = None
        caption_candidates = []
        for sh in text_shapes:
            sl = sh.left / EMU_PER_IN
            st = sh.top / EMU_PER_IN
            txt = shape_text(sh)
            if not txt:
                continue
            if not (left - 0.1 <= sl <= left + 0.4):
                continue
            if not (top - 1.3 <= st <= top - 0.2):
                continue
            if txt in TAG_WORDS:
                tag = txt
            elif txt in MARKER_WORDS:
                continue
            elif is_placeholder(txt):
                continue
            elif is_pri and sh.shape_type == MSO_SHAPE_TYPE.TEXT_BOX and len(txt) < 40 and publisher is None:
                publisher = txt
            else:
                caption_candidates.append(txt)

        caption = max(caption_candidates, key=len) if caption_candidates else ""

        image = None
        best_area = 0.0
        col_center = left + 0.5
        for sh in pictures:
            sl = sh.left / EMU_PER_IN
            st = sh.top / EMU_PER_IN
            sw = sh.width / EMU_PER_IN
            sh_h = sh.height / EMU_PER_IN
            if sw < 0.3 or sh_h < 0.3:
                continue
            if not (top - 1.3 <= st <= top - 0.05):
                continue
            pic_center = sl + sw / 2
            if abs(pic_center - col_center) < 0.45:
                area = sw * sh_h
                if area > best_area:
                    best_area = area
                    image = sh

        posts.append({
            "kind": "pre" if is_pri else "ces",
            "caption": caption,
            "publisher": publisher,
            "metrics": metrics,
            "tag": tag,
            "image": image.image.blob if image is not None else None,
            "image_ext": (image.image.ext or "png").lower() if image is not None else None,
            "_left": left,
            "_top": top,
        })

    posts.sort(key=lambda p: (p["_top"], p["_left"]))
    return posts


def _parse_json_array(raw):
    m = re.search(r"\[.*\]", raw, re.DOTALL)
    if m:
        try:
            return json.loads(m.group(0))
        except json.JSONDecodeError:
            return None
    return None


def _request_descriptions(image_parts, prompt, api_key, model, base_url, temperature, max_tokens):
    content = [{"type": "text", "text": prompt}] + image_parts
    url = base_url.rstrip("/") + "/chat/completions"
    resp = requests.post(
        url,
        headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
        json={"model": model, "messages": [{"role": "user", "content": content}],
              "temperature": temperature, "max_tokens": max_tokens},
        timeout=300,
    )
    resp.raise_for_status()
    return resp.json()


def describe_images(posts, api_key, model, base_url, temperature=0.2, max_tokens=4000, retries=2):
    image_parts = []
    idx_map = []
    for i, p in enumerate(posts):
        if p["image"] is not None:
            image_parts.append(
                {"type": "image_url",
                 "image_url": {"url": image_data_url(p["image"], p.get("image_ext", "png"))}}
            )
            idx_map.append(i)

    if not image_parts:
        return {}, None

    n = len(image_parts)
    base_prompt = (
        f"คุณจะได้รับรูปภาพโพสต์โซเชียลมีเดียจำนวน {n} รูปตามลำดับ: รูปที่ 1 ถึงรูปที่ {n}\n"
        "สำหรับแต่ละรูป ให้บรรยายสิ่งที่เห็นในภาพเป็นภาษาไทย จำนวน 3 ประโยค:\n"
        "ประโยคที่ 1: ฉากหรือเนื้อหาหลักของภาพ บุคคล สิ่งของ สี และการจัดวางองค์ประกอบ\n"
        "ประโยคที่ 2: ข้อความ พาดหัว หรือโลโก้ที่ปรากฏบนภาพ (อ่านข้อความภาษาไทยบนภาพแล้วเขียนออกมาเป็นภาษาไทย)\n"
        "ประโยคที่ 3: สไตล์กราฟิกและโทนหรืออารมณ์โดยรวมของโพสต์\n\n"
        f"ตอบกลับเป็น JSON array จำนวน {n} object โดยมี key: post, image_description เท่านั้น ไม่ต้องมีข้อความอื่น"
    )

    descriptions = None
    data = None
    for attempt in range(retries + 1):
        prompt = base_prompt
        if attempt > 0:
            prompt += "\n\nสำคัญมาก: ตอบกลับเป็น JSON ที่ถูกต้องเท่านั้น ห้ามมีข้อความอื่นนอก JSON"
        data = _request_descriptions(image_parts, prompt, api_key, model, base_url,
                                     temperature, max_tokens)
        raw = data["choices"][0]["message"]["content"]
        descriptions = _parse_json_array(raw)
        if descriptions is not None:
            break

    result = {}
    if descriptions:
        for d in descriptions:
            batch_idx = d.get("post")
            if batch_idx is None:
                continue
            i = batch_idx - 1
            if 0 <= i < len(idx_map):
                result[idx_map[i]] = d.get("image_description")

    return result, data.get("usage") if data else None
