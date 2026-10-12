import io
import os
import random
import time
from concurrent.futures import ThreadPoolExecutor, as_completed

import pandas as pd
import streamlit as st
from dotenv import load_dotenv
from pptx import Presentation

from report_writer.analysis_export import build_analysis_title, export_analysis_excel
from report_writer.ces_data import CES_COLUMNS, load_ces_data, ces_row
from report_writer.config import load_config
from report_writer.key_thematic import BRANDS, build_key_thematic_messages, chat_text
from report_writer.overview_ces import (
    build_overview_ces_messages,
    overview_ces_validation_error,
)
from report_writer.overview_pre import (
    build_overview_pre_messages,
    overview_pre_validation_error,
)
from report_writer.pptx_extractor import describe_images, extract_posts
from report_writer.pre_data import load_pre_data, pre_row
from report_writer.top_pri_content import build_top_pri_messages, group_top_pri_posts
from report_writer.pri_mri import build_pri_mri_messages, load_top_pri_content_workbook, parse_sheet_context, prepare_pri_mri_rows
from report_writer.usage_cost import add_usage, empty_usage, usage_summary

load_dotenv(override=True)

st.set_page_config(page_title="Report Writer", layout="wide")
st.title("Report Writer")

config = load_config()
pricing_by_model = config.get("pricing_usd_per_million_tokens", {})
default_model = config.get("model", "gpt-5.6-luna")
extract_model = config.get("extract_model", default_model)
analysis_model = config.get("analysis_model", default_model)
base_url = config.get("base_url", "https://api.openai.com/v1")
max_completion_tokens = int(config.get("max_completion_tokens", 4000))
analysis_parallel_enabled = bool(config.get("analysis_parallel_enabled", True))
analysis_parallel_workers = max(1, int(config.get("analysis_parallel_workers", 3)))
extract_parallel_enabled = bool(config.get("extract_parallel_enabled", True))
extract_parallel_workers = max(1, int(config.get("extract_parallel_workers", 2)))

api_key = os.getenv("OPENAI_API_KEY", "").strip()

if not api_key:
    st.warning("OPENAI_API_KEY is empty in .env")


@st.cache_data(show_spinner=False)
def load_slides(file_bytes):
    prs = Presentation(io.BytesIO(file_bytes))

    def title_of(slide):
        try:
            t = slide.shapes.title
            if t is not None and t.text.strip():
                return t.text.strip().replace("\t", " ").replace("\n", " ")
        except Exception:
            pass
        return "(no title)"

    titles = {}
    for i, slide in enumerate(prs.slides, start=1):
        titles[i] = title_of(slide)
    return titles


@st.cache_data(show_spinner=False)
def extract_slide(file_bytes, slide_no):
    prs = Presentation(io.BytesIO(file_bytes))
    slide = prs.slides[slide_no - 1]
    posts = extract_posts(slide)
    for p in posts:
        p.pop("_left", None)
        p.pop("_top", None)
    return posts


@st.cache_data(show_spinner=False)
def load_ces(file_bytes):
    df = load_ces_data(file_bytes)
    return df


@st.cache_data(show_spinner=False)
def load_pre(file_bytes):
    return load_pre_data(file_bytes)


@st.cache_data(show_spinner=False)
def load_top_pri_content_sheets(file_bytes):
    return load_top_pri_content_workbook(file_bytes)


def pillar_of(title, pillars):
    for p in pillars:
        if p in title:
            return p
    return "Other"


def brand_of(title):
    for b in ("SCG", "PTT", "CP"):
        if b in title:
            return b
    return None


def kind_of(title):
    if "TOP PRI" in title:
        return "PRE"
    if "TOP POST" in title:
        return "CES"
    return None


def group_label(kind, pillar):
    return f"{kind} {pillar}"


st.header("1. Upload inputs")

pptx_file = st.file_uploader("PPTX file", type=["pptx"], key="pptx_upload")
excel_file = st.file_uploader("Excel file (CES score)", type=["xlsx", "xls", "xlsm"], key="excel_upload")
pre_excel_file = st.file_uploader(
    "Excel file (PRE score summary)",
    type=["xlsx", "xls", "xlsm"],
    key="pre_excel_upload",
)

top_pri_content_file = st.file_uploader(
    "Top PRI Content (for PRI MRI section only)",
    type=["xlsx", "xls", "xlsm"],
    key="top_pri_content_upload",
)


selected_slides = []
titles = {}
ces_df = None
pre_df = None
pillars = config.get("pillars", ["ESG Branding", "Business Branding", "Employer Branding"])

if pptx_file is not None:
    titles = load_slides(pptx_file.getvalue())
    st.caption(f"{len(titles)} slides detected")

    default_slides = [n for n, t in titles.items() if "TOP POST" in t or "TOP PRI" in t]
    selected_slides = st.multiselect(
        "Slides to process",
        options=list(titles.keys()),
        default=default_slides,
        format_func=lambda n: f"{n}: {titles[n]}",
    )

if excel_file is not None:
    try:
        ces_df = load_ces(excel_file.getvalue())
        with st.expander("CES score data (columns reordered: score | growth)"):
            st.dataframe(ces_df, use_container_width=True, height=280)
    except Exception as e:
        st.error(f"Could not read Excel: {e}")

if pre_excel_file is not None:
    try:
        pre_df = load_pre(pre_excel_file.getvalue())
        with st.expander("PRE score data (8 required columns)"):
            st.dataframe(pre_df, use_container_width=True, height=280)
    except Exception as e:
        st.error(f"Could not read PRE Excel: {e}")
if top_pri_content_file is not None:
    try:
        top_pri_content_bytes = top_pri_content_file.getvalue()
        top_pri_content_sheets = load_top_pri_content_sheets(top_pri_content_bytes)
        sheet_names = list(top_pri_content_sheets.keys())
        scg_sheet_names = [sheet_name for sheet_name in sheet_names if "SCG" in str(sheet_name)]
        st.session_state["pri_mri_top_pri_content_sheets"] = top_pri_content_sheets
        st.caption(
            f"Top PRI Content loaded ({len(sheet_names)} sheets, SCG sheets: {len(scg_sheet_names)}) — used only for PRI MRI section"
        )
        with st.expander("Top PRI Content sheets (for PRI MRI)"):
            st.write(pd.DataFrame({"sheet_name": sheet_names, "used_for_pri_mri": ["SCG" in str(sheet_name) for sheet_name in sheet_names]}))
    except Exception as e:
        st.session_state.pop("pri_mri_top_pri_content_sheets", None)
        st.error(f"Could not read Top PRI Content Excel: {e}")
else:
    st.session_state.pop("pri_mri_top_pri_content_sheets", None)

st.header("2. Extract & describe")

if st.button("Run", type="primary", disabled=(api_key == "" or pptx_file is None or not selected_slides)):
    if not api_key:
        st.error("Set OPENAI_API_KEY in .env")
    else:
        groups = {}
        for n in selected_slides:
            kind = kind_of(titles[n]) or "CES"
            p = pillar_of(titles[n], pillars)
            groups.setdefault((kind, p), []).append(n)

        ordered_keys = [(k, p) for k, p in groups if k == "CES"] + [(k, p) for k, p in groups if k != "CES"]

        results = []
        progress = st.progress(0.0, text="Processing slides...")
        total_usage = empty_usage()
        total_slides = len(selected_slides)
        pptx_bytes = pptx_file.getvalue()

        jobs = []
        index = 0
        for kind, pillar in ordered_keys:
            for slide_no in sorted(groups[(kind, pillar)]):
                jobs.append((index, kind, pillar, slide_no))
                index += 1

        def extract_describe_job(job):
            job_index, kind, pillar, slide_no = job
            posts = extract_slide(pptx_bytes, slide_no)
            desc_map, usage = describe_images(
                posts,
                api_key=api_key,
                model=extract_model,
                base_url=base_url,
                max_completion_tokens=max_completion_tokens,
            )
            for post_index, post in enumerate(posts):
                post["image_description"] = desc_map.get(post_index)

            result = {
                "kind": kind,
                "pillar": pillar,
                "group": group_label(kind, pillar),
                "brand": brand_of(titles[slide_no]),
                "slide": slide_no,
                "title": titles[slide_no],
                "posts": posts,
            }
            progress_text = f"{group_label(kind, pillar)}: slide {slide_no}"
            return job_index, result, usage, progress_text

        if extract_parallel_enabled and total_slides > 1:
            max_workers = min(extract_parallel_workers, total_slides)
            indexed_results = []
            with ThreadPoolExecutor(max_workers=max_workers) as executor:
                futures = [executor.submit(extract_describe_job, job) for job in jobs]
                for done, future in enumerate(as_completed(futures), 1):
                    indexed_results.append(future.result())
                    progress.progress(done / total_slides, text=f"Processing slides... {done}/{total_slides}")

            for _, result, usage, _ in sorted(indexed_results, key=lambda item: item[0]):
                results.append(result)
                if usage:
                    add_usage(total_usage, usage)
        else:
            for done, job in enumerate(jobs, 1):
                _, result, usage, progress_text = extract_describe_job(job)
                results.append(result)
                if usage:
                    add_usage(total_usage, usage)
                progress.progress(done / total_slides, text=progress_text)

        progress.empty()
        st.session_state["results"] = results
        st.session_state["usage"] = total_usage
        st.session_state["usage_model"] = extract_model
        if ces_df is not None:
            st.session_state["ces_df"] = ces_df
        if pre_df is not None:
            st.session_state["pre_df"] = pre_df

st.header("3. Results")

results = st.session_state.get("results", [])
if not results:
    st.info("Results will appear here after you run.")
else:
    usage = st.session_state.get("usage", {})
    if usage:
        usage_model = st.session_state.get("usage_model", extract_model)
        st.caption(usage_summary(usage, usage_model, pricing_by_model))

    for r in results:
        with st.expander(f"{r['group']} — Slide {r['slide']} — {r['title']}", expanded=False):
            for p in r["posts"]:
                cols = st.columns([1, 3])
                with cols[0]:
                    if p.get("image") is not None:
                        st.image(p["image"], width=180)
                    else:
                        st.caption("no image")
                with cols[1]:
                    st.markdown(f"**Tag:** {p.get('tag') or '—'}")
                    m = p.get("metrics") or {}
                    if p.get("kind") == "pre":
                        st.caption(
                            f"PRI {m.get('pri', '—'):,} | PRE {m.get('pre', '—')} | "
                            f"Engagement {m.get('engagement', '—'):,}"
                        )
                        if p.get("publisher"):
                            st.caption(f"Source: {p['publisher']}")
                    else:
                        st.caption(
                            f"Engagement {m.get('engagement', '—')} | "
                            f"Like {m.get('like', '—')} | "
                            f"Comment {m.get('comment', '—')} | "
                            f"Share {m.get('share', '—')}"
                        )
                    st.caption(p["caption"][:200])
                    if p.get("image_description"):
                        st.write(p["image_description"])
                    else:
                        st.caption("_no image description_")

st.header("4. Analysis")


def section_instruction(section_name):
    for s in config.get("sections", []):
        if s.get("name") == section_name:
            return s.get("prompt", "")
    return ""


all_results = st.session_state.get("results", [])
ces_results = [result for result in all_results if result.get("kind") == "CES"]
pre_results = [result for result in all_results if result.get("kind") == "PRE"]
has_pre_data = pre_df is not None or st.session_state.get("pre_df") is not None
has_ces_data = ces_df is not None or st.session_state.get("ces_df") is not None
ces_df_now = ces_df if ces_df is not None else st.session_state.get("ces_df")
pre_df_now = pre_df if pre_df is not None else st.session_state.get("pre_df")
pri_mri_sheets_now = st.session_state.get("pri_mri_top_pri_content_sheets", {})
pri_mri_sheet_names = [sheet_name for sheet_name in pri_mri_sheets_now.keys() if "SCG" in str(sheet_name)]

ces_posts_by_combo = {}
for result in ces_results:
    pillar = result.get("pillar")
    brand = result.get("brand")
    if pillar not in pillars or brand not in BRANDS:
        continue
    ces_posts_by_combo.setdefault((pillar, brand), []).extend(result.get("posts") or [])

pre_posts_by_combo = group_top_pri_posts(pre_results, pillars)
key_thematic_combos = [
    (pillar, brand)
    for pillar in pillars
    for brand in BRANDS
    if (pillar, brand) in ces_posts_by_combo
]
top_pri_combos = [
    (pillar, brand)
    for pillar in pillars
    for brand in BRANDS
    if (pillar, brand) in pre_posts_by_combo
]

def _should_retry_analysis(error_text):
    lowered = (error_text or "").lower()
    retry_signals = (
        "429",
        "rate limit",
        "timeout",
        "timed out",
        "temporarily",
        "connection reset",
        "500",
        "502",
        "503",
        "504",
    )
    return any(signal in lowered for signal in retry_signals)


def request_analysis(messages, max_attempts=3):
    for attempt in range(1, max_attempts + 1):
        try:
            output, usage = chat_text(
                messages,
                api_key=api_key,
                model=analysis_model,
                base_url=base_url,
                max_completion_tokens=max_completion_tokens,
            )
            return output, usage, None
        except Exception as e:
            error_text = str(e)
            if attempt >= max_attempts or not _should_retry_analysis(error_text):
                return None, None, error_text
            sleep_seconds = min(1.5 * (2 ** (attempt - 1)), 6.0) + random.uniform(0.0, 0.2)
            time.sleep(sleep_seconds)


def _run_analysis_jobs(jobs, worker, progress_prefix, progress_label):
    section_usage = empty_usage()
    section_results = []
    if not jobs:
        return section_results, section_usage

    progress = st.progress(0.0, text=progress_label)
    total = len(jobs)

    if analysis_parallel_enabled and total > 1:
        indexed_results = []
        max_workers = min(analysis_parallel_workers, total)
        with ThreadPoolExecutor(max_workers=max_workers) as executor:
            futures = [executor.submit(worker, job) for job in jobs]
            for done, future in enumerate(as_completed(futures), 1):
                index, result, usage = future.result()
                indexed_results.append((index, result, usage))
                progress.progress(done / total, text=f"{progress_prefix}: {done}/{total}")

        for _, result, usage in sorted(indexed_results, key=lambda item: item[0]):
            section_results.append(result)
            if usage:
                add_usage(section_usage, usage)
    else:
        for done, job in enumerate(jobs, 1):
            _, result, usage = worker(job)
            section_results.append(result)
            if usage:
                add_usage(section_usage, usage)
            progress.progress(done / total, text=f"{progress_prefix}: {result['pillar']} / {result.get('brand') or 'Overview'}")

    progress.empty()
    return section_results, section_usage
def add_analysis_usage(section_usage):
    total = dict(st.session_state.get("analysis_usage", empty_usage()))
    for key in total:
        value = section_usage.get(key)
        if isinstance(value, int):
            total[key] = total.get(key, 0) + value
    st.session_state["analysis_usage"] = total
    st.session_state["analysis_usage_model"] = analysis_model


def generate_key_thematic(instruction_prompt):
    jobs = list(enumerate(key_thematic_combos))

    def worker(job):
        index, (pillar, brand) = job
        ces_rows = {
            comparison_brand: ces_row(ces_df_now, pillar, comparison_brand)
            for comparison_brand in BRANDS
        }
        messages = build_key_thematic_messages(
            pillar,
            brand,
            ces_rows,
            ces_posts_by_combo[(pillar, brand)],
            instruction=instruction_prompt,
        )
        output, usage, error = request_analysis(messages)
        result = {
            "section": "CES Key Thematic",
            "pillar": pillar,
            "brand": brand,
            "title": build_analysis_title("CES", pillar, brand, "Key Thematic"),
            "output": output,
            "error": error,
            "messages": messages,
        }
        return index, result, usage

    return _run_analysis_jobs(
        jobs,
        worker,
        progress_prefix="Key Thematic",
        progress_label="Generating Key Thematic...",
    )
def generate_top_pri(instruction_prompt):
    jobs = list(enumerate(top_pri_combos))

    def worker(job):
        index, (pillar, brand) = job
        pre_rows = {
            comparison_brand: pre_row(pre_df_now, pillar, comparison_brand)
            for comparison_brand in BRANDS
        }
        messages = build_top_pri_messages(
            pillar,
            brand,
            pre_rows,
            pre_posts_by_combo[(pillar, brand)],
            instruction=instruction_prompt,
        )
        output, usage, error = request_analysis(messages)
        result = {
            "section": "Top PRI Content",
            "pillar": pillar,
            "brand": brand,
            "title": build_analysis_title("PRE", pillar, brand, "Top PRI Content"),
            "output": output,
            "error": error,
            "messages": messages,
        }
        return index, result, usage

    return _run_analysis_jobs(
        jobs,
        worker,
        progress_prefix="Top PRI Content",
        progress_label="Generating Top PRI Content...",
    )
def key_thematic_outputs(section_results):
    return {
        (analysis["pillar"], analysis["brand"]): analysis.get("output")
        for analysis in section_results
        if not analysis.get("error") and analysis.get("output")
    }


def top_pri_outputs(section_results):
    return {
        (analysis["pillar"], analysis["brand"]): analysis.get("output")
        for analysis in section_results
        if not analysis.get("error") and analysis.get("output")
    }


def can_generate_overview_ces(section_results):
    if ces_df_now is None:
        return False
    outputs = key_thematic_outputs(section_results)
    for pillar in pillars:
        ces_rows = {brand: ces_row(ces_df_now, pillar, brand) for brand in BRANDS}
        pillar_outputs = {brand: outputs.get((pillar, brand)) for brand in BRANDS}
        if overview_ces_validation_error(ces_rows, pillar_outputs) is None:
            return True
    return False


def can_generate_overview_pre(section_results):
    if pre_df_now is None:
        return False
    outputs = top_pri_outputs(section_results)
    for pillar in pillars:
        pre_rows = {brand: pre_row(pre_df_now, pillar, brand) for brand in BRANDS}
        pillar_outputs = {brand: outputs.get((pillar, brand)) for brand in BRANDS}
        if overview_pre_validation_error(pre_rows, pillar_outputs) is None:
            return True
    return False


def generate_overview_ces(source_results, instruction_prompt):
    outputs = key_thematic_outputs(source_results)
    section_usage = empty_usage()
    section_results = []
    progress = st.progress(0.0, text="Generating Overview CES...")
    for done, pillar in enumerate(pillars, 1):
        ces_rows = {brand: ces_row(ces_df_now, pillar, brand) for brand in BRANDS}
        pillar_outputs = {brand: outputs.get((pillar, brand)) for brand in BRANDS}
        error = overview_ces_validation_error(ces_rows, pillar_outputs)
        messages = []
        output = None
        usage = None
        if error is None:
            messages = build_overview_ces_messages(
                pillar,
                ces_rows,
                pillar_outputs,
                instruction=instruction_prompt,
            )
            output, usage, error = request_analysis(messages)
        section_results.append({
            "section": "CES Overview",
            "pillar": pillar,
            "brand": None,
            "title": build_analysis_title("CES", pillar, None, "Overview"),
            "output": output,
            "error": error,
            "messages": messages,
        })
        if usage:
            add_usage(section_usage, usage)
        progress.progress(done / len(pillars), text=f"Overview CES: {pillar}")
    progress.empty()
    return section_results, section_usage


def generate_overview_pre(source_results, instruction_prompt):
    outputs = top_pri_outputs(source_results)
    section_usage = empty_usage()
    section_results = []
    progress = st.progress(0.0, text="Generating PRE Overview...")
    for done, pillar in enumerate(pillars, 1):
        pre_rows = {brand: pre_row(pre_df_now, pillar, brand) for brand in BRANDS}
        pillar_outputs = {brand: outputs.get((pillar, brand)) for brand in BRANDS}
        error = overview_pre_validation_error(pre_rows, pillar_outputs)
        messages = []
        output = None
        usage = None
        if error is None:
            messages = build_overview_pre_messages(
                pillar,
                pre_rows,
                pillar_outputs,
                instruction=instruction_prompt,
            )
            output, usage, error = request_analysis(messages)
        section_results.append({
            "section": "PRE Overview",
            "pillar": pillar,
            "brand": None,
            "title": build_analysis_title("PRE", pillar, None, "Overview"),
            "output": output,
            "error": error,
            "messages": messages,
        })
        if usage:
            add_usage(section_usage, usage)
        progress.progress(done / len(pillars), text=f"PRE Overview: {pillar}")
    progress.empty()
    return section_results, section_usage


def generate_pri_mri(instruction_prompt):
    jobs = list(enumerate(pri_mri_sheet_names))

    def worker(job):
        index, sheet_name = job
        sheet_frame = pri_mri_sheets_now.get(sheet_name)
        prepared_rows, error = prepare_pri_mri_rows(sheet_name, sheet_frame)
        brand, pillar = parse_sheet_context(sheet_name)

        messages = []
        output = None
        usage = None
        if error is None:
            messages = build_pri_mri_messages(
                sheet_name,
                prepared_rows,
                instruction=instruction_prompt,
            )
            output, usage, error = request_analysis(messages)

        result = {
            "section": "PRI MRI",
            "pillar": pillar or sheet_name,
            "brand": brand,
            "sheet": sheet_name,
            "title": build_analysis_title("PRE", pillar or sheet_name, brand, "PRI MRI"),
            "output": output,
            "error": error,
            "messages": messages,
        }
        return index, result, usage

    return _run_analysis_jobs(
        jobs,
        worker,
        progress_prefix="PRI MRI",
        progress_label="Generating PRI MRI...",
    )
key_ready = api_key != "" and has_ces_data and bool(key_thematic_combos)
top_pri_ready = api_key != "" and has_pre_data and bool(top_pri_combos)
pri_mri_ready = api_key != "" and bool(pri_mri_sheet_names)

instruction = st.text_area(
    "Key Thematic instruction prompt (editable)",
    value=section_instruction("CES Key Thematic (Enterprise)"),
    key="kt_instruction",
    height=300,
)
if st.button("Analyze Key Thematic", disabled=not key_ready):
    analyses, section_usage = generate_key_thematic(instruction)
    st.session_state["analyses"] = analyses
    st.session_state["overview_ces_analyses"] = []
    add_analysis_usage(section_usage)

top_pri_instruction = st.text_area(
    "Top PRI Content instruction prompt (editable)",
    value=section_instruction("Top PRI Content"),
    key="top_pri_instruction",
    height=300,
)
if st.button("Analyze Top PRI Content", disabled=not top_pri_ready):
    top_pri_analyses, section_usage = generate_top_pri(top_pri_instruction)
    st.session_state["top_pri_analyses"] = top_pri_analyses
    st.session_state["overview_pre_analyses"] = []
    add_analysis_usage(section_usage)

current_key_thematics = st.session_state.get("analyses", [])
overview_ces_instruction = st.text_area(
    "Overview CES instruction prompt (editable)",
    value=section_instruction("CES Overview (Enterprise)"),
    key="overview_ces_instruction",
    height=300,
)
if st.button(
    "Analyze Overview CES",
    disabled=(api_key == "" or not can_generate_overview_ces(current_key_thematics)),
):
    overview_ces_analyses, section_usage = generate_overview_ces(
        current_key_thematics,
        overview_ces_instruction,
    )
    st.session_state["overview_ces_analyses"] = overview_ces_analyses
    add_analysis_usage(section_usage)

current_top_pri = st.session_state.get("top_pri_analyses", [])
overview_pre_instruction = st.text_area(
    "PRE Overview instruction prompt (editable)",
    value=section_instruction("PRE Overview"),
    key="overview_pre_instruction",
    height=300,
)
if st.button(
    "Analyze PRE Overview",
    disabled=(api_key == "" or not can_generate_overview_pre(current_top_pri)),
):
    overview_pre_analyses, section_usage = generate_overview_pre(
        current_top_pri,
        overview_pre_instruction,
    )
    st.session_state["overview_pre_analyses"] = overview_pre_analyses
    add_analysis_usage(section_usage)

pri_mri_instruction = st.text_area(
    "PRI MRI instruction prompt (editable)",
    value=section_instruction("PRI MRI"),
    key="pri_mri_instruction",
    height=300,
)
if st.button("Analyze PRI MRI", disabled=not pri_mri_ready):
    pri_mri_analyses, section_usage = generate_pri_mri(pri_mri_instruction)
    st.session_state["pri_mri_analyses"] = pri_mri_analyses
    add_analysis_usage(section_usage)

analysis_ready = key_ready and top_pri_ready
if st.button("Generate All Analysis", type="primary", disabled=not analysis_ready):
    total_usage = empty_usage()
    analyses, usage = generate_key_thematic(instruction)
    for key in total_usage:
        total_usage[key] += usage.get(key, 0)
    top_pri_analyses, usage = generate_top_pri(top_pri_instruction)
    for key in total_usage:
        total_usage[key] += usage.get(key, 0)
    overview_ces_analyses, usage = generate_overview_ces(
        analyses,
        overview_ces_instruction,
    )
    for key in total_usage:
        total_usage[key] += usage.get(key, 0)
    overview_pre_analyses, usage = generate_overview_pre(
        top_pri_analyses,
        overview_pre_instruction,
    )
    for key in total_usage:
        total_usage[key] += usage.get(key, 0)

    pri_mri_analyses = []
    if pri_mri_ready:
        pri_mri_analyses, usage = generate_pri_mri(pri_mri_instruction)
        for key in total_usage:
            total_usage[key] += usage.get(key, 0)

    st.session_state["analyses"] = analyses
    st.session_state["top_pri_analyses"] = top_pri_analyses
    st.session_state["overview_ces_analyses"] = overview_ces_analyses
    st.session_state["overview_pre_analyses"] = overview_pre_analyses
    st.session_state["pri_mri_analyses"] = pri_mri_analyses
    add_analysis_usage(total_usage)

st.subheader("Analysis results")

analyses = st.session_state.get("analyses", [])
top_pri_analyses = st.session_state.get("top_pri_analyses", [])
overview_ces_analyses = st.session_state.get("overview_ces_analyses", [])
overview_pre_analyses = st.session_state.get("overview_pre_analyses", [])
pri_mri_analyses = st.session_state.get("pri_mri_analyses", [])
combined_analyses = analyses + top_pri_analyses + overview_ces_analyses + overview_pre_analyses + pri_mri_analyses

if not combined_analyses:
    st.info("Analysis results will appear here after generation.")
else:
    analysis_usage = st.session_state.get("analysis_usage", {})
    if analysis_usage:
        analysis_usage_model = st.session_state.get("analysis_usage_model", analysis_model)
        st.caption(usage_summary(analysis_usage, analysis_usage_model, pricing_by_model))

    exportable_analyses = [
        analysis
        for analysis in combined_analyses
        if analysis.get("output") and not analysis.get("error")
    ]
    if exportable_analyses:
        st.download_button(
            "Export analysis to Excel",
            data=export_analysis_excel(exportable_analyses),
            file_name="scg_analysis_results.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    for section_name, section_results in (
        ("Key Thematic", analyses),
        ("Top PRI Content", top_pri_analyses),
        ("Overview CES", overview_ces_analyses),
        ("PRE Overview", overview_pre_analyses),
        ("PRI MRI", pri_mri_analyses),
    ):
        st.markdown(f"**{section_name}**")
        if not section_results:
            st.caption("No results")
            continue
        for analysis in section_results:
            result_label = analysis["pillar"]
            if analysis.get("brand"):
                result_label += f" / {analysis['brand']}"
            with st.expander(result_label, expanded=False):
                if analysis["error"]:
                    st.error(f"Error: {analysis['error']}")
                else:
                    tab_out, tab_data = st.tabs(["Output", "Data sent"])
                    with tab_out:
                        st.code(analysis["output"], language=None, wrap_lines=True)
                    with tab_data:
                        st.code(analysis["messages"][1]["content"], language="text")



